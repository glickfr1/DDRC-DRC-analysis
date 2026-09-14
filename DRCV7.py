import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
from scipy.optimize import curve_fit
from scipy.integrate import trapezoid
import os, re, shutil
from pathlib import Path
from openpyxl import Workbook
from openpyxl.drawing.image import Image
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.utils import get_column_letter

# ---- SETTINGS ----
BASE_DIR = Path(__file__).resolve().parent
input_dir = BASE_DIR / "input"
output_dir = BASE_DIR / "output"

os.makedirs(input_dir, exist_ok=True)
os.makedirs(output_dir, exist_ok=True)
# Output filename is based on the first five characters of the input filename.
# The actual filename is set after the input file is selected.
excel_file = output_dir / "plot_analysis.xlsx"

# ---- SELECT INPUT FILE ----
supported_extensions = {".csv", ".xlsx", ".xls", ".ods"}
input_files = sorted(
    [f for f in input_dir.iterdir() if f.is_file() and f.suffix.lower() in supported_extensions],
    key=lambda f: f.name.lower()
)

if len(input_files) == 0:
    raise RuntimeError(
        f"No supported input files found in {input_dir}. "
        "Supported formats: CSV, Excel (.xlsx/.xls), and OpenDocument (.ods)."
    )

print("\nAvailable input files:\n")

for i, file in enumerate(input_files, 1):
    print(f"{i}. {file.name}")

while True:
    choice = input("\nSelect input file number: ").strip()

    if choice.isdigit():
        idx = int(choice) - 1
        if 0 <= idx < len(input_files):
            input_file = input_files[idx]
            break

    print("Invalid selection. Try again.")

print(f"\nSelected file: {input_file.name}")

# Create a descriptive output filename from the input filename.
input_stem = input_file.stem
prefix = input_stem[:5]
output_name = f"{prefix}_plot_analysis.xlsx"
excel_file = output_dir / output_name

# Avoid overwriting an existing analysis.
if excel_file.exists():
    counter = 2
    while True:
        candidate = output_dir / f"{prefix}_plot_analysis_{counter}.xlsx"
        if not candidate.exists():
            excel_file = candidate
            break
        counter += 1

# ---- REFERENCE COMPOUND ----
ref_input = input(
    "\nReference compound ID for comparison curve (press Enter to skip): "
).strip()
REFERENCE_COMPOUND = ref_input if ref_input else None

if REFERENCE_COMPOUND:
    print(f"Reference compound: {REFERENCE_COMPOUND}")
else:
    print("No reference compound selected.")


# ---- MODEL ----
def four_param_logistic(x, bottom, top, logIC50, hill):
    return bottom + (top-bottom)/(1+10**((logIC50-np.log10(x))*hill))

# ---- WELL PARSER ----
def get_well_col(w):
    if pd.isna(w): return np.nan
    m = re.search(r'[A-Z]+(\d+)', str(w).upper())
    return int(m.group(1)) if m else np.nan

# ---- LOAD ----
suffix = input_file.suffix.lower()

if suffix == ".csv":
    df = pd.read_csv(input_file)
elif suffix in {".xlsx", ".xls"}:
    df = pd.read_excel(input_file)
elif suffix == ".ods":
    df = pd.read_excel(input_file, engine="odf")
else:
    raise RuntimeError(f"Unsupported input format: {suffix}")

# preserve originals for user display
original_columns = list(df.columns)

# lowercase internal copy
df.columns = [c.strip().lower() for c in df.columns]

print("\nDetected columns:\n")

for i, c in enumerate(original_columns):
    print(f"{i+1}. {c}")

print("\nMap the required fields:\n")

def choose_column(field_name):

    while True:

        choice = input(f"Column for '{field_name}': ").strip()

        # allow numeric selection
        if choice.isdigit():

            idx = int(choice)-1

            if 0 <= idx < len(df.columns):
                return df.columns[idx]

        # allow direct typing
        if choice.lower() in df.columns:
            return choice.lower()

        print("Invalid selection. Try again.")

plate_col = choose_column("plate")
well_col_name = choose_column("well")
compound_col = choose_column("compound id")
conc_col = choose_column("concentration")
lum_col = choose_column("raw signal / response")

# optional batch column
batch_input = input(
    "Column for batch (press Enter to skip): "
).strip()

batch_col = None

if batch_input:

    if batch_input.isdigit():
        idx = int(batch_input)-1

        if 0 <= idx < len(df.columns):
            batch_col = df.columns[idx]

    elif batch_input.lower() in df.columns:
        batch_col = batch_input.lower()

# rename internally
rename_map = {
    plate_col: "plate",
    well_col_name: "well",
    compound_col: "compound id",
    conc_col: "concentration",
    lum_col: "raw_signal"
}

if batch_col:
    rename_map[batch_col] = "batch"

df = df.rename(columns=rename_map)

# clean
df["compound id"] = df["compound id"].astype(str).str.strip()
df["well"] = df["well"].astype(str).str.upper().str.strip()
df["well_col"] = df["well"].apply(get_well_col)

# ---- CONTROL SETUP + NORMALIZATION ----
def parse_items(text):
    return [item.strip().upper() for item in re.split(r"[,;\s]+", text) if item.strip()]

print("\nASSAY CONTROL SETUP")
print("1. Controls are arranged in columns")
print("2. Controls are arranged in rows")
print("3. Controls are specified as individual wells")
print("4. No controls are present")

while True:
    control_mode = input("Select control arrangement (1-4): ").strip()
    if control_mode in {"1", "2", "3", "4"}:
        break
    print("Invalid selection. Try again.")

if control_mode == "1":
    neg_cols = {int(x) for x in parse_items(input("Negative-control column(s), e.g. 23 or 23,24: ")) if x.isdigit()}
    pos_cols = {int(x) for x in parse_items(input("Positive-control column(s), e.g. 24 or 24,25: ")) if x.isdigit()}
    control_selector = lambda w, vals: get_well_col(w) in vals
elif control_mode == "2":
    neg_rows = set(parse_items(input("Negative-control row(s), e.g. A or A,B: ")))
    pos_rows = set(parse_items(input("Positive-control row(s), e.g. P or O,P: ")))
    control_selector = lambda w, vals: str(w).strip().upper()[:1] in vals
elif control_mode == "3":
    neg_wells = set(parse_items(input("Negative-control wells, e.g. A1,A2,A3: ")))
    pos_wells = set(parse_items(input("Positive-control wells, e.g. P22,P23,P24: ")))
    control_selector = lambda w, vals: str(w).strip().upper() in vals

zprime_results = []

if control_mode == "4":
    print("\nNo assay controls selected.")
    print("1. Input response is already Percent inhibition (0-100)")
    print("2. Analyze raw signal directly")
    while True:
        response_mode = input("Select response type (1-2): ").strip()
        if response_mode in {"1", "2"}:
            break
        print("Invalid selection. Try again.")
    df["response"] = pd.to_numeric(df["raw_signal"], errors="coerce")
    response_label = "Percent inhibition (%)" if response_mode == "1" else "Raw signal"
    normalized = response_mode == "1"
else:
    df["response"] = np.nan
    for plate, pdf in df.groupby("plate"):
        neg = pdf[pdf["well"].map(lambda w: control_selector(w, neg_cols if control_mode == "1" else neg_rows if control_mode == "2" else neg_wells))]["raw_signal"]
        pos = pdf[pdf["well"].map(lambda w: control_selector(w, pos_cols if control_mode == "1" else pos_rows if control_mode == "2" else pos_wells))]["raw_signal"]
        if len(neg) == 0 or len(pos) == 0:
            print(f"WARNING: Plate {plate} has no usable negative and/or positive controls; normalization skipped for this plate.")
            continue
        nm, pm = neg.mean(), pos.mean()
        nsd, psd = neg.std(ddof=1), pos.std(ddof=1)
        if len(neg) >= 3 and len(pos) >= 3 and abs(nm-pm) >= 1e-6:
            zprime_results.append({"plate": plate, "Z_prime": 1 - (3*(nsd+psd)/abs(nm-pm))})
        if abs(nm-pm) < 1e-6:
            print(f"WARNING: Plate {plate} has nearly identical positive and negative controls; normalization skipped.")
            continue
        mask = df["plate"] == plate
        df.loc[mask, "response"] = 100*(nm - df.loc[mask, "raw_signal"])/(nm-pm)
    response_label = "Percent inhibition (%)"
    normalized = True

df["response"] = df["response"].clip(0, 100) if normalized else df["response"]
df["npi"] = df["response"]

# ---- REMOVE CONTROLS ----
df_fit = df[
    df["compound id"].notna() &
    (df["compound id"] != "") &
    (~df["compound id"].str.lower().isin(["nan", "none"]))
]

# ---- REFERENCE ----
ref_df = df_fit[df_fit["compound id"] == REFERENCE_COMPOUND]
df_fit = df_fit[df_fit["compound id"] != REFERENCE_COMPOUND]

ref_summary = None

if REFERENCE_COMPOUND and ref_df.empty:
    print(
        f"WARNING: Reference compound '{REFERENCE_COMPOUND}' was not found "
        "in the input data. Continuing without a reference curve."
    )

if not ref_df.empty:

    ref_summary = ref_df.groupby("concentration").agg(
        npi_mean=("npi", "mean"),
        npi_sem=("npi", lambda x: x.std(ddof=1)/np.sqrt(len(x)))
    ).reset_index()

# ---- METRICS ----
def compute_metrics(x, y, popt):

    yfit = four_param_logistic(x, *popt)

    rmse = np.sqrt(np.mean((y-yfit)**2))
    dyn = np.max(y)-np.min(y)+1e-6
    norm_rmse = rmse/dyn

    diffs = np.diff(y)

    mono = np.mean(
        np.sign(diffs) == np.sign(np.mean(diffs))
    ) if len(diffs) > 0 else 0

    amp = abs(popt[1]-popt[0])
    hill = abs(popt[3])

    return yfit, norm_rmse, mono, amp, hill

# ⭐ NEW: Plateau scoring
def plateau_score(x, y):

    if len(x) < 6:
        return 0.0

    idx = np.argsort(x)

    x = x[idx]
    y = y[idx]

    top_y = y[:3]
    bot_y = y[-3:]

    def flatness(vals):
        return 1 - (np.std(vals) / (np.abs(np.mean(vals)) + 1e-6))

    def closeness(vals, target):
        return 1 - (np.abs(np.mean(vals) - target) / 100)

    top_score = (
        0.5 * flatness(top_y) +
        0.5 * closeness(top_y, np.max(y))
    )

    bot_score = (
        0.5 * flatness(bot_y) +
        0.5 * closeness(bot_y, np.min(y))
    )

    return max(0, (top_score + bot_score) / 2)

def classify(amp, mono, rmse, hill):

    if amp >= 80 and mono > 0.85 and rmse < 0.15 and 0.3 < hill < 5:
        return "1.1"

    if amp >= 80 and mono > 0.7:
        return "1.2"

    if amp >= 50 and mono > 0.7:
        return "1.3"

    if amp >= 50:
        return "2.1"

    if amp >= 20:
        return "2.2"

    if amp > 10:
        return "3"

    return "4"

def confidence(amp, rmse, mono, n):

    return round(
        min(amp/100,1)*35 +
        max(0,1-rmse)*35 +
        mono*20 +
        min(n/10,1)*10,
        1
    )

def flags(amp, rmse, mono, n, x=None, y=None):

    f = []

    if amp < 20:
        f.append("weak")

    if rmse > 0.25:
        f.append("poor_fit")

    if mono < 0.6:
        f.append("noisy")

    if n < 5:
        f.append("few_points")

    if x is not None and y is not None:

        if plateau_score(x, y) < 0.4:
            f.append("no_plateau")

    return ";".join(f)

def ddrc_curve_score(x, y, hill, chi2):

    # ---- SPAN SCORE (0-30) ----
    span = np.max(y) - np.min(y)

    if span >= 90:
        span_score = 30
    elif span >= 70:
        span_score = 25
    elif span >= 50:
        span_score = 18
    elif span >= 30:
        span_score = 10
    else:
        span_score = 3

    # ---- HILL SCORE (0-15) ----
    ahill = abs(hill)

    if 0.8 <= ahill <= 1.5:
        hill_score = 15
    elif 0.5 <= ahill <= 2.5:
        hill_score = 12
    elif 0.3 <= ahill <= 4:
        hill_score = 8
    else:
        hill_score = 2

    # ---- CHI2 SCORE (0-25) ----
    if chi2 < 50:
        chi_score = 25
    elif chi2 < 150:
        chi_score = 20
    elif chi2 < 400:
        chi_score = 12
    elif chi2 < 1000:
        chi_score = 5
    else:
        chi_score = 0

    # ---- PLATEAU SCORE (0-30) ----
    idx = np.argsort(x)
    y_sorted = y[idx]

    top = y_sorted[:3]
    bottom = y_sorted[-3:]

    def plateau_quality(vals):

        spread = np.std(vals)

        if spread < 3:
            return 15
        elif spread < 7:
            return 10
        elif spread < 12:
            return 5
        else:
            return 0

    plateau_score_val = (
        plateau_quality(top) +
        plateau_quality(bottom)
    )

    total = (
        span_score +
        hill_score +
        chi_score +
        plateau_score_val
    )

    return round(min(total, 100), 1)

# ---- HEATMAP ----
def make_plate_heatmap(df_plate, plate_name, outdir):

    grid = np.full((16,24), np.nan)

    for _, row in df_plate.iterrows():

        m = re.match(r"([A-P])(\d+)", str(row["well"]))

        if not m:
            continue

        r = ord(m.group(1)) - ord('A')
        c = int(m.group(2)) - 1

        if 0 <= r < 16 and 0 <= c < 24:
            grid[r,c] = row["npi"]

    plt.figure(figsize=(10,6))

    plt.imshow(grid, aspect='auto')

    plt.colorbar(label="Percent inhibition (%)")

    plt.title(plate_name)

    fname = os.path.join(outdir, f"{plate_name}_heatmap.png")

    plt.savefig(fname, dpi=200)

    plt.close()

    return fname

# ---- FIT ----
plot_dir = os.path.join(output_dir, "plots")

os.makedirs(plot_dir, exist_ok=True)

results = []

for cmpd, grp in df_fit.groupby("compound id"):

    batch = grp["batch"].iloc[0] if "batch" in grp else ""

    summary = grp.groupby("concentration").agg(
        npi_mean=("npi","mean"),
        npi_sem=("npi", lambda x: x.std(ddof=1)/np.sqrt(len(x)) if len(x) > 1 else np.nan),
        n_replicates=("npi", "count")
    ).reset_index()

    rep_min = int(summary["n_replicates"].min())
    rep_max = int(summary["n_replicates"].max())
    replicate_label = str(rep_min) if rep_min == rep_max else f"{rep_min}-{rep_max}"

    x = summary["concentration"].values
    y = summary["npi_mean"].values
    yerr = summary["npi_sem"].values

    mask = (x > 0) & (~np.isnan(y))

    x, y, yerr = x[mask], y[mask], yerr[mask]

    if len(x) < 4:
        continue

    try:

        p0 = [
            min(y),
            max(y),
            np.log10(np.median(x)),
            1
        ]

        popt, _ = curve_fit(
            four_param_logistic,
            x,
            y,
            p0,
            maxfev=10000
        )

        IC50 = 10**popt[2]

        yfit, rmse, mono, amp, hill = compute_metrics(x, y, popt)

        cls = classify(amp, mono, rmse, hill)
        conf = confidence(amp, rmse, mono, len(x))
        flg = flags(amp, rmse, mono, len(x), x, y)

        xfit = np.logspace(
            np.log10(min(x)),
            np.log10(max(x)),
            100
        )

        yfit_full = four_param_logistic(xfit, *popt)

        auc = trapezoid(yfit_full, np.log10(xfit))

        chi2 = np.sum(
            (y - four_param_logistic(x, *popt))**2
        )

        ddrc_score = ddrc_curve_score(x, y, hill, chi2)

        # ---- PLOT ----
        plt.figure()

        plt.errorbar(
            x,
            y,
            yerr=yerr,
            fmt='o',
            capsize=3
        )

        plt.plot(
            xfit,
            yfit_full,
            label=f"{cmpd} IC50={IC50:.3g}"
        )

        if ref_summary is not None:

            rx = ref_summary["concentration"].values
            ry = ref_summary["npi_mean"].values

            mask_r = rx > 0

            if sum(mask_r) >= 4:

                r_popt, _ = curve_fit(
                    four_param_logistic,
                    rx[mask_r],
                    ry[mask_r],
                    [
                        min(ry),
                        max(ry),
                        np.log10(np.median(rx)),
                        1
                    ],
                    maxfev=10000
                )

                rxfit = np.logspace(
                    np.log10(min(rx)),
                    np.log10(max(rx)),
                    100
                )

                ryfit = four_param_logistic(rxfit, *r_popt)

                plt.plot(
                    rxfit,
                    ryfit,
                    linestyle=":",
                    linewidth=2,
                    label=f"Ref ({REFERENCE_COMPOUND})"
                )

        plt.xscale("log")
        plt.xlabel("Concentration (µM)")
        plt.ylabel(response_label)
        plt.ylim(0,100)

        plt.legend()

        plt.tight_layout()

        plot_path = os.path.join(plot_dir, f"{cmpd}.png")

        plt.savefig(plot_path, dpi=200)

        plt.close()

        results.append({
            "Compound": cmpd,
            "Batch": batch,
            "Fit status": "OK",
            "Replicates": replicate_label,
            "IC50": IC50,
            "Chi-square": chi2,
            "AUC": auc,
            "Curve class": cls,
            "Fit confidence score": conf,
            "DDRC Curve Score": ddrc_score,
            "Flags": flg,
            "Plot": plot_path
        })

    except Exception as exc:
        results.append({"Compound": cmpd, "Batch": batch, "Fit status": f"Fit failed: {type(exc).__name__}: {exc}", "Replicates": replicate_label, "Flags": "fit_failed"})
        continue

if len(results) == 0:
    raise RuntimeError("No successful fits")

# ---- HEATMAPS ----
heatmap_paths = []

for plate, pdf in df.groupby("plate"):

    p = make_plate_heatmap(pdf, plate, plot_dir)

    heatmap_paths.append((plate, p))

# ---- EXCEL ----
wb = Workbook()

ws = wb.active
ws.title = "Results"

headers = ["Compound", "Batch", "Fit status", "Replicates", "IC50", "Chi-square", "AUC", "Curve class", "Fit confidence score", "DDRC Curve Score", "Flags", "Plot"]

ws.append(headers)

for i, r in enumerate(results, 2):

    ws.append([
        r.get(h, "") if h != "Plot" else ""
        for h in headers
    ])

    if os.path.exists(r["Plot"]):

        img = Image(r["Plot"])

        img.width = 300
        img.height = 200

        plot_col = headers.index("Plot") + 1
        plot_cell = f"{get_column_letter(plot_col)}{i}"
        ws.add_image(img, plot_cell)

        ws.row_dimensions[i].height = 150

# Z'
ws_z = wb.create_sheet("Zprime")

ws_z.append(["Plate", "Z_prime"])

for z in zprime_results:
    ws_z.append([z["plate"], z["Z_prime"]])

# IC50 heatmap
ws_h = wb.create_sheet("IC50 Heatmap")

heat_df = pd.DataFrame(results)[["Compound", "IC50"]]

for _, row in heat_df.iterrows():
    ws_h.append(row.tolist())

rule = ColorScaleRule(
    start_type='min',
    end_type='max'
)

ws_h.conditional_formatting.add(
    f"B2:B{len(heat_df)+1}",
    rule
)

# Plate heatmaps
ws_p = wb.create_sheet("Plate Maps")

row = 1

for plate, path in heatmap_paths:

    ws_p.cell(row=row, column=1).value = plate

    img = Image(path)

    img.width = 500
    img.height = 300

    ws_p.add_image(img, f"A{row+1}")

    row += 20

wb.save(excel_file)

shutil.rmtree(plot_dir)

print(f"\nDone → {excel_file}")